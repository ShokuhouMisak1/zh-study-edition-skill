#!/usr/bin/env python3
"""Extract the text of a source PDF / slide deck into per-page files, for reading and chunking.

Usage:
  python extract_source.py <source.pdf|.pptx|.key> <out_dir>

Writes:
  <out_dir>/full.txt          whole document (pdftotext -layout), with '=== page N ===' markers
  <out_dir>/pages/pNNN.txt    one file per page (for slides: one file per slide)
  <out_dir>/outline.txt       PDF bookmarks if any, plus lines that look like headings
                              ("3.2 Basic Feasible Directions", "Chapter 4", "Index of definitions", ...)
Prints page count and the outline so you can plan chapters / lectures.

.pptx is converted with LibreOffice (soffice --headless --convert-to pdf) when available; otherwise
python-pptx text is used (no page images then). Needs poppler (pdftotext, pdfinfo).
Formulas in extracted text are often garbled: always check formulas against the rendered page
(pdftoppm -r 110 -f N -l N -png src.pdf out).
"""
import pathlib
import re
import shutil
import subprocess
import sys


def run(*cmd):
    return subprocess.run(cmd, capture_output=True, text=True)


def pptx_to_pdf(src: pathlib.Path, out: pathlib.Path) -> pathlib.Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    run(soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(src))
    pdf = out / (src.stem + ".pdf")
    return pdf if pdf.exists() else None


def pptx_text(src: pathlib.Path, out: pathlib.Path):
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("need LibreOffice (soffice) or `pip install python-pptx` to read .pptx")
    prs = Presentation(str(src))
    pages = out / "pages"
    pages.mkdir(parents=True, exist_ok=True)
    full = []
    for i, slide in enumerate(prs.slides, 1):
        txt = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                txt.append(sh.text_frame.text)
        if slide.has_notes_slide:
            txt.append("[notes] " + slide.notes_slide.notes_text_frame.text)
        t = "\n".join(txt)
        (pages / f"p{i:03d}.txt").write_text(t, encoding="utf-8")
        full.append(f"=== page {i} ===\n{t}")
    (out / "full.txt").write_text("\n".join(full), encoding="utf-8")
    print(f"{len(prs.slides)} slides (text only, via python-pptx)")


HEAD = re.compile(r"^\s*((chapter|lecture|part|appendix)\s+\w+|\d+(\.\d+){0,2}\s+[A-Z][\w\- ,:'’()]+$|index of .+|exercises|references|bibliography)\s*$", re.I)


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src = pathlib.Path(sys.argv[1]).expanduser()
    out = pathlib.Path(sys.argv[2]).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    if src.suffix.lower() in (".pptx", ".ppt", ".key"):
        pdf = pptx_to_pdf(src, out)
        if pdf is None:
            return pptx_text(src, out)
        src = pdf
    info = run("pdfinfo", str(src)).stdout
    n = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    pages = out / "pages"
    pages.mkdir(exist_ok=True)
    full, heads = [], []
    for i in range(1, n + 1):
        t = run("pdftotext", "-layout", "-f", str(i), "-l", str(i), str(src), "-").stdout
        (pages / f"p{i:03d}.txt").write_text(t, encoding="utf-8")
        full.append(f"=== page {i} ===\n{t}")
        for line in t.splitlines():
            if HEAD.match(line) and len(line.strip()) < 90:
                heads.append(f"p{i:>4}  {line.strip()}")
    (out / "full.txt").write_text("\n".join(full), encoding="utf-8")
    bm = run("pdfinfo", "-listbookmarks", str(src)).stdout if "listbookmarks" in run("pdfinfo", "-h").stderr else ""
    text = (("# PDF bookmarks\n" + bm + "\n") if bm.strip() else "") + "# heading-like lines (page  text)\n" + "\n".join(heads)
    (out / "outline.txt").write_text(text, encoding="utf-8")
    print(f"{n} pages -> {out}/pages/, full text -> {out}/full.txt")
    print(text[:6000])


if __name__ == "__main__":
    main()
